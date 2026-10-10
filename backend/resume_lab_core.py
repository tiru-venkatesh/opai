"""Resume Lab core: pure functions (no FastAPI, no DB, no network, no LLM) so every rule is unit-testable.

Principles:
  * Everything is derived from text the student wrote. Nothing is invented: no number, tool, employer or skill is added.
  * Reviews and coaching are deterministic rules; they say what to check, never what to claim.
  * A tailored resume is the student's own resume with bullets re-ordered / removed, plus bullets copied verbatim from
    their other versions. Edits to wording are the student's, never the system's.
"""
from __future__ import annotations

import hashlib
import io
import re
from collections import Counter
from typing import Any, Dict, List, Optional, Tuple
from xml.sax.saxutils import escape as _xml

_BULLET_START = re.compile(r"^\s*(?:[-*•–●▪>]|\d+[.)])\s+")
_SECTION = re.compile(r"^\s*(education|experience|work experience|projects?|skills|technical skills|achievements|"
                      r"certifications?|summary|objective|publications?|research|positions? of responsibility|"
                      r"extra[- ]?curriculars?|awards?)\s*:?\s*$", re.I)
_METRIC = re.compile(r"(\d+(?:\.\d+)?\s?%|\b\d[\d,]*\+?\s?(?:k|m|x|ms|users|requests|queries|docs|documents|students|records|"
                     r"samples|models|apis?|endpoints|tests|hours|days|reps|stars)\b|\$\s?\d|\b\d{2,}\b)", re.I)
_VERBS = set("""built developed designed implemented created led engineered optimized reduced improved increased automated
deployed trained fine-tuned finetuned integrated architected analyzed researched wrote launched migrated refactored
scaled shipped evaluated benchmarked managed organized authored published contributed collaborated mentored
proposed prototyped debugged tested containerized orchestrated extracted classified detected generated""".split())
_WEAK = ["responsible for", "worked on", "helped", "involved in", "assisted", "various", "etc", "team player",
         "hard working", "hardworking", "duties included", "familiar with"]
_STOP = set("""the and for with you our your will are this that from have has not but all any can able work working team
role intern internship company looking candidate strong good must skills experience years year including etc
who what when where into over such their they them also more than using use used new""".split())
_TERM_RE = re.compile(r"[a-z][a-z0-9+#.\-]{1,}")


def _h(text: str) -> str:
    return hashlib.sha256(re.sub(r"\s+", " ", (text or "").strip().lower()).encode()).hexdigest()


def _contains(text_low: str, term: str) -> bool:
    return re.search(r"(?<![a-z0-9+#])" + re.escape(term) + r"(?![a-z0-9+#])", text_low) is not None


# ---------------------------------------------------------------- parsing
def parse_resume(text: str) -> Dict[str, Any]:
    """Splits a resume into sections and bullets. Pure text processing."""
    sections: Dict[str, List[str]] = {}
    bullets: List[Dict[str, str]] = []
    cur = "header"
    for raw in (text or "").splitlines():
        ln = raw.strip()
        if not ln:
            continue
        m = _SECTION.match(ln)
        if m:
            cur = m.group(1).lower()
            sections.setdefault(cur, [])
            continue
        sections.setdefault(cur, []).append(ln)
        is_b = bool(_BULLET_START.match(ln))
        body = _BULLET_START.sub("", ln).strip()
        first = (body.split() or [""])[0].lower().strip(",.:;")
        if cur in ("skills", "technical skills", "header", "education"):
            continue
        if is_b or (len(body) >= 45 and first in _VERBS):
            if len(body) >= 12:
                bullets.append({"text": body, "section": cur})
    return {"sections": sections, "bullets": bullets}


def bullet_stats(bullets: List[Dict[str, str]]) -> Dict[str, Any]:
    n = len(bullets)
    if not n:
        return {"bullets": 0, "metric_ratio": 0.0, "verb_ratio": 0.0, "avg_words": 0.0, "weak": {}}
    metric = sum(1 for b in bullets if _METRIC.search(b["text"]))
    verb = sum(1 for b in bullets if (b["text"].split() or [""])[0].lower().strip(",.:;") in _VERBS)
    weak = Counter()
    for b in bullets:
        low = b["text"].lower()
        for w in _WEAK:
            if w in low:
                weak[w] += 1
    return {"bullets": n, "metric_ratio": round(metric / n, 2), "verb_ratio": round(verb / n, 2),
            "avg_words": round(sum(len(b["text"].split()) for b in bullets) / n, 1), "weak": dict(weak)}


# =====================================================================================================
# 1. ATS-readiness review (rules on the plain text; no real ATS is simulated)
# =====================================================================================================
_EMAIL = re.compile(r"[\w.+\-]+@[\w\-]+\.[\w.\-]+")
_PHONE = re.compile(r"(?:\+?\d[\s\-()]?){9,14}\d")
_LINK = re.compile(r"(?:linkedin\.com|github\.com|https?://|www\.)", re.I)
_YEAR = re.compile(r"\b(?:19|20)\d{2}\b|\bpresent\b|\bcurrent\b", re.I)
_FANCY = re.compile(r"[★☆✔✓✗✦❖➤➢►◆◇■□✪❑❒✱✲]")
_SKILL_SPLIT = re.compile(r"[,;|/•]")


def ats_review(text: str, terms: Optional[List[str]] = None) -> Dict[str, Any]:
    parsed = parse_resume(text)
    secs = {k for k in parsed["sections"] if k != "header"}
    st = bullet_stats(parsed["bullets"])
    words = len((text or "").split())
    lines = [l for l in (text or "").splitlines() if l.strip()]
    checks: List[Dict[str, str]] = []

    def add(cid: str, label: str, status: str, detail: str, fix: str = "") -> None:
        checks.append({"id": cid, "label": label, "status": status, "detail": detail, "fix": fix})

    # contact
    add("email", "Email address", "ok" if _EMAIL.search(text or "") else "bad",
        "Found." if _EMAIL.search(text or "") else "No email address found.", "" if _EMAIL.search(text or "") else "Put a plain-text email near the top.")
    add("phone", "Phone number", "ok" if _PHONE.search(text or "") else "warn",
        "Found." if _PHONE.search(text or "") else "No phone number found.", "" if _PHONE.search(text or "") else "Add a phone number recruiters can call.")
    add("link", "LinkedIn / GitHub / portfolio link", "ok" if _LINK.search(text or "") else "warn",
        "Found." if _LINK.search(text or "") else "No profile or portfolio link found.", "" if _LINK.search(text or "") else "Add your GitHub or LinkedIn URL as plain text.")
    # sections
    has_work = bool(secs & {"experience", "work experience", "projects", "project", "research"})
    add("work", "Experience or Projects section", "ok" if has_work else "bad",
        "Present." if has_work else "No Experience / Projects heading detected.", "" if has_work else "Use a standard heading such as Projects or Experience.")
    add("edu", "Education section", "ok" if "education" in secs else "warn",
        "Present." if "education" in secs else "No Education heading detected.", "" if "education" in secs else "Add an Education heading with degree, college and years.")
    has_sk = bool(secs & {"skills", "technical skills"})
    add("skills", "Skills section", "ok" if has_sk else "warn",
        "Present." if has_sk else "No Skills heading detected.", "" if has_sk else "Add a Skills section listing tools you really use.")
    # length
    if words > 750:
        add("length", "Length", "warn", f"{words} words (about {round(words / 500, 1)} pages).", "Students usually fit on one page. Cut older or weaker bullets.")
    elif words < 150:
        add("length", "Length", "warn", f"Only {words} words.", "Add detail to projects: what you built, with what, and the result.")
    else:
        add("length", "Length", "ok", f"{words} words.")
    # bullets
    n = st["bullets"]
    if n == 0:
        add("bullets", "Bullet points", "bad", "No bullet points detected.", "Describe each project or role in short bullets starting with a verb.")
    else:
        add("verbs", "Bullets start with action verbs", "ok" if st["verb_ratio"] >= 0.6 else "warn",
            f"{round(100 * st['verb_ratio'])}% of bullets.", "" if st["verb_ratio"] >= 0.6 else "Lead with what you did: Built, Implemented, Reduced...")
        add("metrics", "Bullets with real numbers", "ok" if st["metric_ratio"] >= 0.4 else "warn",
            f"{round(100 * st['metric_ratio'])}% of bullets.", "" if st["metric_ratio"] >= 0.4 else "Add figures only where true: users, documents, % change, time saved.")
        weak = st["weak"]
        add("weak", "Vague phrases", "warn" if weak else "ok", ("Found: " + ", ".join(f'"{w}"' for w in weak)) if weak else "None found.",
            "Replace with the specific action and result." if weak else "")
        longb = [b["text"] for b in parsed["bullets"] if len(b["text"].split()) > 35]
        add("long", "Overlong bullets", "warn" if longb else "ok", f"{len(longb)} bullet(s) over 35 words." if longb else "None.",
            "Split long bullets in two." if longb else "")
    # parsing risks visible in text
    tabbed = [l for l in lines if l.count("\t") >= 2 or l.count("|") >= 3]
    add("columns", "Tables / multi-column layout", "warn" if tabbed else "ok",
        f"{len(tabbed)} line(s) look like table or column content." if tabbed else "Reads as single column.",
        "ATS parsers often scramble columns and tables. Use one column." if tabbed else "")
    fancy = _FANCY.findall(text or "")
    add("symbols", "Decorative symbols", "warn" if fancy else "ok", f"{len(fancy)} decorative symbol(s)." if fancy else "None.",
        "Use plain - or • bullets; some parsers drop other symbols." if fancy else "")
    veryl = [l for l in lines if len(l) > 220]
    add("longlines", "Very long lines", "warn" if veryl else "ok", f"{len(veryl)} very long line(s)." if veryl else "None.",
        "Break long lines; they may be merged table cells." if veryl else "")
    add("dates", "Dates on experience", "ok" if _YEAR.search(text or "") else "warn",
        "Years found." if _YEAR.search(text or "") else "No years or date ranges found.", "" if _YEAR.search(text or "") else "Add dates to education and roles.")

    keyword: Optional[Dict[str, Any]] = None
    if terms:
        low = (text or "").lower()
        have = [t for t in terms if _contains(low, t)]
        cov = len(have) / len(terms)
        keyword = {"coverage": round(cov, 2), "present": have, "missing": [t for t in terms if t not in have],
                   "note": "Add a missing term only if it is true for you."}
        add("keywords", "Keywords from the posting", "ok" if cov >= 0.6 else "warn" if cov >= 0.3 else "bad",
            f"{len(have)} of {len(terms)} posting terms appear.", "" if cov >= 0.6 else "Mention the tools you genuinely used that the posting asks for.")

    bad = sum(1 for c in checks if c["status"] == "bad")
    warn = sum(1 for c in checks if c["status"] == "warn")
    score = max(0, 100 - 15 * bad - 6 * warn)
    return {"score": score, "grade": "Strong" if score >= 85 else "Okay" if score >= 65 else "Needs work", "checks": checks,
            "keywords": keyword, "words": words,
            "method": "Rule-based checks on your resume text. No real ATS is simulated; different systems parse differently."}


# =====================================================================================================
# 2. Bullet coach (diagnose + ask; a rewrite is only accepted if it adds nothing new)
# =====================================================================================================
_WEAK_STARTS = ("responsible for", "worked on", "helped", "involved in", "assisted", "participated in", "tasked with",
                "duties included", "worked with")
_PASSIVE = re.compile(r"\b(?:was|were|been|being|is|are)\s+\w+ed\b", re.I)
_VERB_OPTIONS = ["Built", "Implemented", "Designed", "Developed", "Automated", "Reduced", "Improved", "Analyzed", "Deployed"]
_NUM = re.compile(r"\d[\d,]*(?:\.\d+)?")


def _strip_bullet(t: str) -> str:
    return _BULLET_START.sub("", (t or "").strip()).strip()


def coach_bullet(text: str, terms: Optional[List[str]] = None) -> Dict[str, Any]:
    b = _strip_bullet(text)
    low = b.lower()
    words = b.split()
    first = (words[0].lower().strip(",.:;") if words else "")
    issues: List[Dict[str, str]] = []
    questions: List[str] = []

    weak_start = next((w for w in _WEAK_STARTS if low.startswith(w)), None)
    if weak_start:
        issues.append({"id": "weak_start", "severity": "high", "message": f'Starts with "{weak_start}", which hides what you actually did.'})
        questions.append("What exactly did you do? Name the action (built, wrote, tested, deployed...).")
    elif first not in _VERBS and not first.endswith("ed"):
        issues.append({"id": "no_verb", "severity": "medium", "message": "Does not start with an action verb."})
    if not _METRIC.search(b):
        issues.append({"id": "no_metric", "severity": "medium", "message": "No number or outcome."})
        questions.append("Is there a real figure you can add? For example how many users, documents or tests, a % change, or time saved. Leave it out if you do not know it.")
    if len(words) > 28:
        issues.append({"id": "too_long", "severity": "medium", "message": f"{len(words)} words; aim for one clear line (10-26 words)."})
    elif 0 < len(words) < 8:
        issues.append({"id": "too_short", "severity": "medium", "message": "Very short; add what you used and what happened."})
        questions.append("Which tools or methods did you use?")
    if _PASSIVE.search(b):
        issues.append({"id": "passive", "severity": "low", "message": "Passive wording; say what you did, not what was done."})
    for w in _WEAK:
        if w in low and w != weak_start and w not in ("helped", "assisted"):
            issues.append({"id": "vague", "severity": "low", "message": f'Vague phrase: "{w}".'})
            break
    in_bullet, optional = [], []
    if terms:
        in_bullet = [t for t in terms if _contains(low, t)]
        optional = [t for t in terms if t not in in_bullet][:4]
    return {"bullet": b, "words": len(words), "issues": issues, "questions": questions,
            "verb_options": _VERB_OPTIONS if (weak_start or first not in _VERBS) else [],
            "skeleton": "<Action verb> <what you built or did> using <tools you used> to <result, with a real number if you have one>",
            "posting_terms_already_covered": in_bullet,
            "posting_terms_only_if_true": optional,
            "ok": not any(i["severity"] in ("high", "medium") for i in issues)}


def _caps(s: str) -> set:
    toks = re.findall(r"\b[A-Z][A-Za-z0-9+#.\-]{1,}\b", s or "")
    return {t.lower() for t in toks}


def rewrite_is_grounded(original: str, rewrite: str, facts: str = "", known_terms: Optional[List[str]] = None) -> Tuple[bool, List[str]]:
    """A suggested rewrite is acceptable only if it adds no number, tool or proper noun that the student did not provide."""
    base = f"{original} {facts}"
    probs: List[str] = []
    new_nums = {n.replace(",", "") for n in _NUM.findall(rewrite)} - {n.replace(",", "") for n in _NUM.findall(base)}
    if new_nums:
        probs.append("adds numbers not in your text: " + ", ".join(sorted(new_nums)))
    base_low, rw_low = base.lower(), (rewrite or "").lower()
    new_terms = [t for t in (known_terms or []) if _contains(rw_low, t) and not _contains(base_low, t)]
    if new_terms:
        probs.append("adds tools/skills you did not mention: " + ", ".join(new_terms))
    rw_first = (rewrite.split() or [""])[0].lower()
    new_caps = {c for c in _caps(rewrite) - _caps(base) if c != rw_first}
    if new_caps:
        probs.append("adds names/terms not in your text: " + ", ".join(sorted(new_caps)))
    return (not probs, probs)


# =====================================================================================================
# 3. Compare two versions
# =====================================================================================================
def _norm_b(t: str) -> str:
    return re.sub(r"[^a-z0-9 ]+", "", re.sub(r"\s+", " ", t.lower())).strip()


def _jacc(a: str, b: str) -> float:
    x, y = set(_norm_b(a).split()), set(_norm_b(b).split())
    return len(x & y) / len(x | y) if (x | y) else 0.0


_SKILL_LINE = re.compile(r"^\s*(?:technical\s+)?(?:skills?|languages|tools|frameworks|libraries|technologies)\s*:\s*(.+)$", re.I)


def _skill_set(parsed: Dict[str, Any]) -> set:
    """Skills come from a `Skills: a, b, c` line (anywhere) or from lines under a Skills heading."""
    out = set()
    for sec, lns in parsed["sections"].items():
        for ln in lns:
            m = _SKILL_LINE.match(ln)
            body = m.group(1) if m else (ln if sec in ("skills", "technical skills") else None)
            if body:
                out |= {x.strip().lower() for x in _SKILL_SPLIT.split(body) if 1 < len(x.strip()) <= 30}
    return out


def compare_versions(text_a: str, text_b: str, terms: Optional[List[str]] = None) -> Dict[str, Any]:
    pa, pb = parse_resume(text_a), parse_resume(text_b)
    ba, bb = [b["text"] for b in pa["bullets"]], [b["text"] for b in pb["bullets"]]
    ha, hb = {_norm_b(t): t for t in ba}, {_norm_b(t): t for t in bb}
    common = [ha[k] for k in ha if k in hb]
    only_a = [ha[k] for k in ha if k not in hb]
    only_b = [hb[k] for k in hb if k not in ha]
    reworded, used = [], set()
    for a in list(only_a):
        best, bj = None, 0.0
        for j, b in enumerate(only_b):
            if j in used:
                continue
            s = _jacc(a, b)
            if s > bj:
                best, bj = j, s
        if best is not None and bj >= 0.55:
            used.add(best)
            reworded.append({"a": a, "b": only_b[best], "similarity": round(bj, 2)})
    ra = {r["a"] for r in reworded}
    rb = {r["b"] for r in reworded}
    sa, sb = bullet_stats(pa["bullets"]), bullet_stats(pb["bullets"])
    ska, skb = _skill_set(pa), _skill_set(pb)
    out: Dict[str, Any] = {
        "common_bullets": common, "only_in_a": [t for t in only_a if t not in ra], "only_in_b": [t for t in only_b if t not in rb],
        "reworded": reworded,
        "skills": {"only_in_a": sorted(ska - skb), "only_in_b": sorted(skb - ska), "shared": sorted(ska & skb)},
        "sections": {"only_in_a": sorted(set(pa["sections"]) - set(pb["sections"]) - {"header"}),
                     "only_in_b": sorted(set(pb["sections"]) - set(pa["sections"]) - {"header"})},
        "stats": {"a": {**sa, "words": len(text_a.split())}, "b": {**sb, "words": len(text_b.split())},
                  "delta_b_minus_a": {"metric_ratio": round(sb["metric_ratio"] - sa["metric_ratio"], 2),
                                      "verb_ratio": round(sb["verb_ratio"] - sa["verb_ratio"], 2),
                                      "bullets": sb["bullets"] - sa["bullets"]}}}
    if terms:
        la, lb = text_a.lower(), text_b.lower()
        pa_t = [t for t in terms if _contains(la, t)]
        pb_t = [t for t in terms if _contains(lb, t)]
        out["posting"] = {"a": {"coverage": round(len(pa_t) / len(terms), 2), "missing": [t for t in terms if t not in pa_t]},
                          "b": {"coverage": round(len(pb_t) / len(terms), 2), "missing": [t for t in terms if t not in pb_t]}}
    return out


# =====================================================================================================
# 4. Tailored build (reorder / drop the student's bullets; add verbatim bullets from their other versions)
# =====================================================================================================
_NO_BULLET_SECTIONS = ("skills", "technical skills", "header", "education")


def _is_bullet(line: str, section: str) -> bool:
    if section in _NO_BULLET_SECTIONS:
        return False
    body = _BULLET_START.sub("", line).strip()
    first = (body.split() or [""])[0].lower().strip(",.:;")
    return bool(_BULLET_START.match(line) or (len(body) >= 45 and first in _VERBS)) and len(body) >= 12


def build_tailored(base_text: str, exclude: Optional[List[str]] = None, extras: Optional[List[Dict[str, str]]] = None,
                   order_terms: Optional[List[str]] = None, max_per_group: Optional[int] = None) -> Dict[str, Any]:
    """Returns {sections:[{title, items:[{kind:'text'|'bullet', text, origin}]}], dropped, added}. Never alters wording."""
    excl = {_norm_b(t) for t in (exclude or [])}
    sections: List[Dict[str, Any]] = [{"title": None, "key": "header", "items": []}]
    cur = sections[0]
    for raw in (base_text or "").splitlines():
        ln = raw.strip()
        if not ln:
            continue
        m = _SECTION.match(ln)
        if m:
            cur = {"title": ln.rstrip(":").strip(), "key": m.group(1).lower(), "items": []}
            sections.append(cur)
            continue
        if _is_bullet(ln, cur["key"]):
            cur["items"].append({"kind": "bullet", "text": _BULLET_START.sub("", ln).strip(), "origin": "base"})
        else:
            cur["items"].append({"kind": "text", "text": ln, "origin": "base"})

    dropped: List[str] = []
    for s in sections:
        kept = []
        for it in s["items"]:
            if it["kind"] == "bullet" and _norm_b(it["text"]) in excl:
                dropped.append(it["text"])
            else:
                kept.append(it)
        s["items"] = kept

    # reorder bullets inside each group (a group = consecutive bullets under the same non-bullet line)
    def hits(t: str) -> int:
        low = t.lower()
        return sum(1 for x in (order_terms or []) if _contains(low, x))

    if order_terms or max_per_group:
        for s in sections:
            out, run = [], []
            def flush():
                nonlocal run
                if run:
                    ordered = sorted(run, key=lambda it: -hits(it["text"])) if order_terms else run
                    out.extend(ordered[:max_per_group] if max_per_group else ordered)
                    run = []
            for it in s["items"]:
                if it["kind"] == "bullet":
                    run.append(it)
                else:
                    flush()
                    out.append(it)
            flush()
            s["items"] = out

    added: List[str] = []
    have = {_norm_b(it["text"]) for s in sections for it in s["items"] if it["kind"] == "bullet"}
    for ex in (extras or []):
        t = _strip_bullet(ex.get("text", ""))
        if not t or _norm_b(t) in have or _norm_b(t) in excl:   # an excluded bullet stays out even if re-suggested
            continue
        want = (ex.get("section") or "").lower()
        target = next((s for s in sections if s["key"] == want and s["key"] not in _NO_BULLET_SECTIONS), None) \
            or next((s for s in reversed(sections) if any(i["kind"] == "bullet" for i in s["items"])), None)
        if target is None:
            target = {"title": "Projects", "key": "projects", "items": []}
            sections.append(target)
        last = max((i for i, it in enumerate(target["items"]) if it["kind"] == "bullet"), default=len(target["items"]) - 1)
        target["items"].insert(last + 1, {"kind": "bullet", "text": t, "origin": "other_version"})
        have.add(_norm_b(t))
        added.append(t)
    return {"sections": sections, "dropped": dropped, "added": added}


def doc_bullets(doc: Dict[str, Any]) -> List[str]:
    return [it["text"] for s in doc["sections"] for it in s["items"] if it["kind"] == "bullet"]


def render_txt(doc: Dict[str, Any]) -> str:
    out: List[str] = []
    for s in doc["sections"]:
        if s["title"]:
            out += ["", s["title"].upper()]
        for it in s["items"]:
            out.append(("- " if it["kind"] == "bullet" else "") + it["text"])
    return "\n".join(out).strip() + "\n"


def _head(doc: Dict[str, Any]) -> Tuple[str, List[str]]:
    items = [it["text"] for it in doc["sections"][0]["items"]]
    return (items[0] if items else "Resume"), items[1:]


def render_docx(doc: Dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt
    d = Document()
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Inches(0.7)
        sec.top_margin = sec.bottom_margin = Inches(0.6)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    name, contact = _head(doc)
    p = d.add_paragraph()
    r = p.add_run(name)
    r.bold, r.font.size = True, Pt(16)
    for c in contact:
        d.add_paragraph(c).paragraph_format.space_after = Pt(0)
    for s in doc["sections"][1:]:
        h = d.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        hr = h.add_run((s["title"] or "").upper())
        hr.bold = True
        for it in s["items"]:
            if it["kind"] == "bullet":
                d.add_paragraph(it["text"], style="List Bullet").paragraph_format.space_after = Pt(0)
            else:
                para = d.add_paragraph(it["text"])
                para.paragraph_format.space_after = Pt(0)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_pdf(doc: Dict[str, Any]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
                            title=_head(doc)[0], author=_head(doc)[0])
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=10, leading=13)
    h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=16, leading=19)
    h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=7, spaceAfter=2)
    bl = ParagraphStyle("bl", parent=base, leftIndent=12, bulletIndent=2)
    name, contact = _head(doc)
    flow: list = [Paragraph(_xml(name), h1)] + [Paragraph(_xml(c), base) for c in contact]
    for s in doc["sections"][1:]:
        flow.append(Paragraph(_xml((s["title"] or "").upper()), h2))
        for it in s["items"]:
            flow.append(Paragraph(_xml(it["text"]), bl, bulletText="\u2022") if it["kind"] == "bullet" else Paragraph(_xml(it["text"]), base))
    flow.append(Spacer(1, 2))
    pdf.build(flow)
    return buf.getvalue()


RENDERERS = {"txt": ("text/plain; charset=utf-8", lambda d: render_txt(d).encode("utf-8")),
             "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", render_docx),
             "pdf": ("application/pdf", render_pdf)}
