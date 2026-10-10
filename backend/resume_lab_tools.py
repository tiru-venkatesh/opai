"""Pure text tools for Resume Lab: parsing, review, coaching, comparison, build validation and export rendering.

No FastAPI, no database, no network in this module, so every function is deterministic and unit-testable
(tests/test_resume_lab_tools.py runs without any server). resume_lab_api.py is a thin layer on top.

Principle kept from the rest of Resume Lab: nothing here invents experience. Review and coaching only
measure and point at what is written; builds only select the user's own bullets or bullets the user typed
and marked as theirs; rewrites from a model are verified against the source text before they are shown.
"""
from __future__ import annotations

import hashlib
import io
import re
from collections import Counter
from typing import Any, Dict, List, Optional, Tuple

_BULLET_START = re.compile(r"^\s*(?:[\u2022\u25e6\u25aa\u25cf\u25cb\u2023]\s*|[-*\u2013>]\s+|\d+[.)]\s+)")
_SECTION = re.compile(r"^\s*(education|experience|work experience|relevant experience|internship experience|internships?|projects?|"
                      r"academic projects?|skills|technical skills|achievements|certifications?|summary|professional summary|"
                      r"objective|publications?|research|positions? of responsibility|extra[- ]?curriculars?|awards?|"
                      r"coding profiles?|coding marks|leadership|activities)\s*:?\s*$", re.I)
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_PHONE = re.compile(r"(?<!\d)(?:\+?\d[\s-]?){9,13}(?!\d)")
_URL = re.compile(r"(?:https?://|www\.)\S+|\b(?:linkedin|github)\.com/\S*", re.I)
_METRIC = re.compile(r"(\d+(?:\.\d+)?\s?%|\b\d[\d,]*\+?\s?(?:k|m|x|ms|users|requests|queries|docs|documents|students|records|"
                     r"samples|models|apis?|endpoints|tests|hours|days|reps|stars)\b|\$\s?\d|\b\d{2,}\b)", re.I)
_VERBS = set("""built developed designed implemented created led engineered optimized reduced improved increased automated
deployed trained fine-tuned finetuned integrated architected analyzed researched wrote launched migrated refactored
scaled shipped evaluated benchmarked managed organized authored published contributed collaborated mentored
proposed prototyped debugged tested containerized orchestrated extracted classified detected generated""".split())
_WEAK = ["responsible for", "worked on", "helped", "involved in", "assisted", "various", "etc", "team player",
         "hard working", "hardworking", "duties included", "familiar with"]
_STOP = set("""the and for with you our your will are this that from have has not but all any can able work working team
role intern internship company looking candidate strong good must skills experience years year including etc
who what when where into over such their they them also more than using use used new""".split())
_TERM_RE = re.compile(r"[a-z][a-z0-9+#.\-]{1,}")




def _h(text: str) -> str:
    return hashlib.sha256(re.sub(r"\s+", " ", (text or "").strip().lower()).encode()).hexdigest()


# ---------------------------------------------------------------- parsing
def clean_pdf_text(text: str) -> str:
    """Repairs common PDF-extraction damage: letter-spaced lines ("P y T o r c h") and stray form feeds."""
    out = []
    for ln in (text or "").replace("\x0c", "\n").splitlines():
        toks = ln.strip().split(" ")
        singles = sum(1 for t in toks if len(t) == 1)
        if len(ln.strip()) >= 5 and len(toks) >= 4 and singles / len(toks) > 0.6:
            words = re.split(r" {2,}", ln.strip())
            ln = " ".join(w.replace(" ", "") for w in words)
        out.append(ln.rstrip())
    return "\n".join(out)


def scrub_pii(text: str) -> str:
    """Reference resumes belong to other people: drop emails, phone numbers and profile links before storing."""
    t = _EMAIL.sub("[email]", text or "")
    t = _URL.sub("[link]", t)
    return _PHONE.sub("[phone]", t)


_CONTENT = ("experience", "work experience", "relevant experience", "internship experience", "internship", "internships",
            "projects", "project", "academic projects", "academic project", "achievements", "research", "leadership",
            "positions of responsibility", "activities", "publications", "publication")


_DATE_RANGE = re.compile(r"\b(?:19|20)\d\d\b\s*(?:-|\u2013|\u2014|to)\s*(?:(?:19|20)\d\d|present|current|ongoing)\b|\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+(?:19|20)\d\d\b", re.I)


def _looks_like_subheading(body: str) -> bool:
    """A project/role title line ('CivicIQ | React, Node.js', 'Intern, Acme  Jun 2025 - Aug 2025') must never be glued
    onto the bullet above it as if it were a wrapped line."""
    if "|" in body or _DATE_RANGE.search(body):
        return True
    w = body.split()
    return 0 < len(w) <= 6 and sum(1 for x in w if x[:1].isupper()) >= max(2, len(w) - 1)


def parse_resume(text: str) -> Dict[str, Any]:
    """Splits a resume into sections and bullets. Pure text processing; tolerant of PDF text."""
    sections: Dict[str, List[str]] = {}
    bullets: List[Dict[str, str]] = []
    cur = "header"
    last: Optional[int] = None          # index of the bullet that wrapped lines may continue
    for raw in clean_pdf_text(text).splitlines():
        ln = raw.strip()
        if not ln:
            continue
        head = ln.rstrip(":").strip()
        m = _SECTION.match(ln)
        if m or (head.isupper() and 3 <= len(head) <= 40 and len(head.split()) <= 4 and not re.search(r"\d", head)):
            cur = (m.group(1) if m else head).lower()
            sections.setdefault(cur, [])
            last = None
            continue
        sections.setdefault(cur, []).append(ln)
        if cur not in _CONTENT:
            last = None
            continue
        is_b = bool(_BULLET_START.match(ln))
        body = _BULLET_START.sub("", ln).strip()
        first = (body.split() or [""])[0].lower().strip(",.:;")
        if is_b or first in _VERBS:
            if len(body) >= 12:
                bullets.append({"text": body, "section": cur})
                last = len(bullets) - 1
            continue
        if last is not None and (body[:1].islower() or (not bullets[last]["text"].rstrip().endswith(".") and not _looks_like_subheading(body))):
            bullets[last]["text"] += " " + body          # wrapped continuation
        elif len(body) >= 60:
            bullets.append({"text": body, "section": cur})  # statement written without a bullet glyph
            last = len(bullets) - 1
        else:
            last = None
    return {"sections": sections, "bullets": bullets}


def bullet_stats(bullets: List[Dict[str, str]]) -> Dict[str, Any]:
    n = len(bullets)
    if not n:
        return {"bullets": 0, "metric_ratio": 0.0, "verb_ratio": 0.0, "avg_words": 0.0, "weak": {}}
    metric = sum(1 for b in bullets if _METRIC.search(b["text"]))
    verb = sum(1 for b in bullets if (b["text"].split() or [""])[0].lower().strip(",.:;") in _VERBS)
    weak = Counter()
    for b in bullets:
        low = b["text"].lower()
        for w in _WEAK:
            if w in low:
                weak[w] += 1
    return {"bullets": n, "metric_ratio": round(metric / n, 2), "verb_ratio": round(verb / n, 2),
            "avg_words": round(sum(len(b["text"].split()) for b in bullets) / n, 1), "weak": dict(weak)}




# =====================================================================================================
# Shared helpers
# =====================================================================================================
def contains_term(text_low: str, term: str) -> bool:
    return re.search(r"(?<![a-z0-9+#])" + re.escape(term) + r"(?![a-z0-9+#])", text_low) is not None


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9+#%.]+", " ", (s or "").lower()).strip()


def _words(s: str) -> set:
    return {w for w in _norm(s).split() if len(w) > 1}


def similarity(a: str, b: str) -> float:
    """Token Jaccard on normalised words. 1.0 = same words, order ignored."""
    wa, wb = _words(a), _words(b)
    if not wa or not wb:
        return 0.0
    return len(wa & wb) / len(wa | wb)


def _first_word(text: str) -> str:
    return (text.split() or [""])[0].lower().strip(",.:;")


# =====================================================================================================
# Structure: raw text <-> ordered sections (what the builder edits and the exporters render)
# =====================================================================================================
def _title(heading: str) -> str:
    h = heading.strip().rstrip(":").strip()
    return h.title() if h.isupper() or h.islower() else h


def to_structure(text: str) -> Dict[str, Any]:
    """Ordered structure of a resume: {"header": [lines], "sections": [{"key","name","items":[{"type","text"}]}]}.
    type is 'bullet' (an achievement line) or 'line' (a sub-heading such as 'Project | Stack' or a plain line).
    Uses the same heading/bullet rules as parse_resume so both views agree."""
    header: List[str] = []
    sections: List[Dict[str, Any]] = []
    cur: Optional[Dict[str, Any]] = None
    last: Optional[Dict[str, Any]] = None
    for raw in clean_pdf_text(text).splitlines():
        ln = raw.strip()
        if not ln:
            continue
        head = ln.rstrip(":").strip()
        m = _SECTION.match(ln)
        if m or (head.isupper() and 3 <= len(head) <= 40 and len(head.split()) <= 4 and not re.search(r"\d", head)):
            key = (m.group(1) if m else head).lower()
            cur = {"key": key, "name": _title(head), "items": []}
            sections.append(cur)
            last = None
            continue
        if cur is None:
            header.append(ln)
            continue
        is_b = bool(_BULLET_START.match(ln))
        body = _BULLET_START.sub("", ln).strip()
        if cur["key"] not in _CONTENT:
            cur["items"].append({"type": "bullet" if is_b else "line", "text": body})
            last = None
            continue
        first = _first_word(body)
        if is_b or first in _VERBS:
            if len(body) >= 12:
                last = {"type": "bullet", "text": body}
                cur["items"].append(last)
            else:
                cur["items"].append({"type": "line", "text": body})
                last = None
            continue
        if last is not None and (body[:1].islower() or (not last["text"].rstrip().endswith(".") and not _looks_like_subheading(body))):
            last["text"] += " " + body
        elif len(body) >= 60:
            last = {"type": "bullet", "text": body}
            cur["items"].append(last)
        else:
            cur["items"].append({"type": "line", "text": body})
            last = None
    return {"header": header, "sections": sections}


def render_text(struct: Dict[str, Any]) -> str:
    out: List[str] = list(struct.get("header") or [])
    if struct.get("summary"):
        out += ["", "SUMMARY", struct["summary"]]
    for s in struct.get("sections") or []:
        out += ["", (s.get("name") or "").upper()]
        for it in s.get("items") or []:
            out.append(("- " if it.get("type") == "bullet" else "") + it.get("text", ""))
    return "\n".join(out).strip() + "\n"


# =====================================================================================================
# 1. Review: ATS / quality checks on extracted text
# =====================================================================================================
_STD_SECTIONS = {"education", "experience", "work experience", "relevant experience", "internship experience", "internship", "internships",
                 "projects", "project", "academic projects", "academic project", "skills", "technical skills", "achievements",
                 "certifications", "certification", "summary", "professional summary", "objective", "publications", "publication",
                 "research", "positions of responsibility", "position of responsibility", "extracurriculars", "extra-curriculars",
                 "extra curriculars", "awards", "award", "coding profiles", "coding marks", "leadership", "activities"}
_ALIASES = {"education": {"education"}, "experience_or_projects": {"experience", "work experience", "relevant experience", "internship experience",
            "internship", "internships", "projects", "project", "academic projects", "academic project", "research"}, "skills": {"skills", "technical skills"}}
_OK_NON_ASCII = set("•◦▪●○‣–—‘’“”…·°×→é")
_MONTHS = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)"
_DATE_STYLES = {"Mon YYYY": re.compile(r"\b" + _MONTHS + r"\.?\s+(?:19|20)\d\d\b", re.I),
                "Month YYYY": re.compile(r"\b(?:january|february|march|april|june|july|august|september|october|november|december)\s+(?:19|20)\d\d\b", re.I),
                "MM/YYYY": re.compile(r"\b(?:0?[1-9]|1[0-2])/(?:19|20)\d\d\b"),
                "YYYY-MM": re.compile(r"\b(?:19|20)\d\d-(?:0[1-9]|1[0-2])\b")}
_FIRST_PERSON = re.compile(r"(^|\s)(i|my|me|myself)(\s|[,.]|$)", re.I)
_BUZZ = ["passionate", "synergy", "go-getter", "self-motivated", "results-oriented", "detail-oriented", "dynamic", "rockstar", "ninja", "guru"]
_VAGUE = ["various", "several", "many", "some", "etc", "stuff", "things", "a lot of", "different"]


def _chk(cid: str, area: str, status: str, message: str, fix: str = "", evidence: Optional[List[str]] = None) -> Dict[str, Any]:
    return {"id": cid, "area": area, "status": status, "message": message, "fix": fix, "evidence": (evidence or [])[:5]}


def review_resume(text: str, terms: Optional[List[str]] = None) -> Dict[str, Any]:
    """Deterministic review of the text a parser would see. No model, no invented advice: every finding names what it measured."""
    parsed = parse_resume(text)
    bl = parsed["bullets"]
    keys = {k for k in parsed["sections"] if k != "header"}
    words = len(text.split())
    checks: List[Dict[str, Any]] = []

    # --- structure
    missing = [lab for lab, al in (("Education", _ALIASES["education"]), ("Experience or Projects", _ALIASES["experience_or_projects"]),
                                   ("Skills", _ALIASES["skills"])) if not (keys & al)]
    if not missing:
        checks.append(_chk("core_sections", "structure", "pass", "Education, Experience/Projects and Skills sections are all present."))
    else:
        st = "fail" if "Experience or Projects" in missing else "warn"
        checks.append(_chk("core_sections", "structure", st, "Missing section(s): " + ", ".join(missing) + ".",
                           "Add them under standard headings so parsers file your content correctly."))
    odd = sorted(k for k in keys if k not in _STD_SECTIONS)
    checks.append(_chk("standard_headings", "structure", "warn" if odd else "pass",
                       ("Non-standard heading(s): " + ", ".join(odd) + ".") if odd else "All headings are standard.",
                       "Rename to common headings such as Education, Projects, Experience, Skills." if odd else "", odd))
    head_txt = " ".join(parsed["sections"].get("header", [])[:8]) or text[:400]
    has_email, has_phone = bool(_EMAIL.search(head_txt)), bool(_PHONE.search(head_txt))
    has_link = bool(re.search(r"github\.com|linkedin\.com|portfolio|https?://", text, re.I))
    miss = [n for n, ok in (("email", has_email), ("phone", has_phone)) if not ok]
    checks.append(_chk("contact", "structure", "fail" if miss else "pass",
                       ("Contact line is missing: " + ", ".join(miss) + ".") if miss else "Email and phone are in the header.",
                       "Put email and phone as plain text at the top, not inside an image or text box." if miss else ""))
    checks.append(_chk("profile_links", "structure", "pass" if has_link else "warn",
                       "A GitHub / LinkedIn / portfolio link is present." if has_link else "No GitHub, LinkedIn or portfolio link found.",
                       "" if has_link else "Add your GitHub and LinkedIn as plain text URLs."))

    # --- formatting (text level only)
    lt = "pass" if words <= 650 else "warn" if words <= 900 else "fail"
    checks.append(_chk("length", "formatting", lt, f"About {words} words (roughly {max(1, round(words / 550 + 0.2))} page(s)).",
                       "" if lt == "pass" else "Internship screens expect one page. Cut the weakest bullets first."))
    long_lines = [ln.strip()[:80] for ln in text.splitlines() if len(ln.strip()) > 170]
    checks.append(_chk("line_length", "formatting", "warn" if long_lines else "pass",
                       f"{len(long_lines)} very long line(s); columns or tables may have been merged." if long_lines else "No merged-column lines detected.",
                       "Use a single column. Avoid tables and text boxes." if long_lines else "", long_lines))
    spaced = [ln.strip()[:60] for ln in text.splitlines() if len(ln.strip()) >= 5 and len(ln.strip().split(" ")) >= 4
              and sum(1 for t in ln.strip().split(" ") if len(t) == 1) / len(ln.strip().split(" ")) > 0.6]
    checks.append(_chk("spaced_text", "formatting", "warn" if spaced else "pass",
                       "Some text extracts letter-by-letter (e.g. 'P y T o r c h')." if spaced else "Text extracts as normal words.",
                       "Re-export the PDF from a text editor with real fonts; avoid decorative letter-spacing." if spaced else "", spaced))
    odd_chars = sorted({ch for ch in text if ord(ch) > 127 and ch not in _OK_NON_ASCII and not ch.isalpha()})
    checks.append(_chk("special_characters", "formatting", "warn" if odd_chars else "pass",
                       ("Unusual symbols that parsers often drop: " + " ".join(odd_chars[:8])) if odd_chars else "No unusual symbols.",
                       "Replace icons and special glyphs with plain words." if odd_chars else ""))
    glyphs = {m.group(0).strip() for ln in text.splitlines() if (m := re.match(r"^\s*([\u2022\u25e6\u25aa\u25cf\u25cb\u2023]|[-*\u2013>])", ln))}
    checks.append(_chk("bullet_glyphs", "formatting", "warn" if len(glyphs) > 1 else "pass",
                       ("Mixed bullet styles: " + " ".join(sorted(glyphs))) if len(glyphs) > 1 else "Bullet style is consistent.",
                       "Use one bullet style throughout." if len(glyphs) > 1 else ""))
    styles = [n for n, rx in _DATE_STYLES.items() if rx.search(text)]
    checks.append(_chk("date_format", "formatting", "warn" if len(styles) > 1 else "pass",
                       ("Dates are written in different styles: " + ", ".join(styles)) if len(styles) > 1 else "Date format is consistent.",
                       "Pick one format (for example 'Jun 2025') and use it everywhere." if len(styles) > 1 else ""))

    # --- content
    n = len(bl)
    st_ = bullet_stats(bl)
    if n < 4:
        checks.append(_chk("bullet_count", "content", "fail", f"Only {n} achievement bullet(s) found.",
                           "Add 2-4 bullets per project or role describing what you did and the result."))
    else:
        checks.append(_chk("bullet_count", "content", "warn" if n > 18 else "pass", f"{n} bullets found.", "Keep to the strongest 12-16." if n > 18 else ""))
    if n:
        mr, vr = st_["metric_ratio"], st_["verb_ratio"]
        checks.append(_chk("metrics", "content", "pass" if mr >= 0.5 else "warn" if mr >= 0.3 else "fail",
                           f"{round(100 * mr)}% of bullets contain a number or measurable outcome.",
                           "" if mr >= 0.5 else "Add real figures (users, records, % change, time saved) where you have them. Never estimate to look better.",
                           [b["text"][:90] for b in bl if not _METRIC.search(b["text"])]))
        checks.append(_chk("action_verbs", "content", "pass" if vr >= 0.7 else "warn" if vr >= 0.4 else "fail",
                           f"{round(100 * vr)}% of bullets start with an action verb.",
                           "" if vr >= 0.7 else "Start each bullet with what you did (Built, Designed, Reduced...).",
                           [b["text"][:90] for b in bl if _first_word(b["text"]) not in _VERBS]))
        weak = [b["text"][:90] for b in bl if any(w in b["text"].lower() for w in _WEAK)]
        checks.append(_chk("weak_phrases", "content", "warn" if weak else "pass",
                           f"{len(weak)} bullet(s) use filler phrases ({', '.join(sorted(st_['weak']))})." if weak else "No filler phrases.",
                           "Replace with the specific action and result." if weak else "", weak))
        longb = [b["text"][:90] for b in bl if len(b["text"].split()) > 30]
        shortb = [b["text"][:90] for b in bl if len(b["text"].split()) < 7]
        checks.append(_chk("bullet_length", "content", "warn" if (longb or shortb) else "pass",
                           f"{len(longb)} bullet(s) over 30 words, {len(shortb)} under 7." if (longb or shortb) else "Bullet lengths are readable.",
                           "Aim for one line each, about 10-26 words." if (longb or shortb) else "", longb + shortb))
        openers = Counter(_first_word(b["text"]) for b in bl)
        rep = [f"{w} x{c}" for w, c in openers.most_common(3) if c >= 3 and w]
        checks.append(_chk("repeated_openers", "content", "warn" if rep else "pass",
                           ("Repeated opening words: " + ", ".join(rep)) if rep else "Bullet openers vary.",
                           "Vary the verbs you lead with." if rep else ""))
        fp = [b["text"][:90] for b in bl if _FIRST_PERSON.search(b["text"])]
        buzz = sorted({w for w in _BUZZ if w in text.lower()})
        checks.append(_chk("tone", "content", "warn" if (fp or buzz) else "pass",
                           ("First-person wording in " + str(len(fp)) + " bullet(s). " if fp else "") + (("Buzzwords: " + ", ".join(buzz)) if buzz else "")
                           if (fp or buzz) else "Tone is direct.", "Drop 'I/my' and replace buzzwords with evidence." if (fp or buzz) else "", fp))

    # --- keywords (only when a posting was supplied)
    if terms:
        low = text.lower()
        have = [t for t in terms if contains_term(low, t)]
        miss_t = [t for t in terms if t not in have]
        cov = len(have) / len(terms)
        checks.append(_chk("keyword_coverage", "keywords", "pass" if cov >= 0.6 else "warn" if cov >= 0.35 else "fail",
                           f"{len(have)} of {len(terms)} posting terms appear in this resume ({round(100 * cov)}%).",
                           "" if not miss_t else "Missing: " + ", ".join(miss_t[:10]) + ". Add a term only if it is true for you.", miss_t))
        stuffed = [f"{t} x{len(re.findall(r'(?<![a-z0-9+#])' + re.escape(t) + r'(?![a-z0-9+#])', low))}" for t in terms
                   if len(re.findall(r"(?<![a-z0-9+#])" + re.escape(t) + r"(?![a-z0-9+#])", low)) >= 6]
        checks.append(_chk("keyword_stuffing", "keywords", "warn" if stuffed else "pass",
                           ("Repeated unusually often: " + ", ".join(stuffed)) if stuffed else "No keyword stuffing.",
                           "Mention a skill where you used it, not repeatedly." if stuffed else ""))

    weights = {"pass": 1.0, "warn": 0.5, "fail": 0.0}
    scores: Dict[str, int] = {}
    for area in ("structure", "formatting", "content", "keywords"):
        cs = [c for c in checks if c["area"] == area]
        if cs:
            scores[area] = round(100 * sum(weights[c["status"]] for c in cs) / len(cs))
    overall = round(sum(scores.values()) / len(scores)) if scores else 0
    order = {"fail": 0, "warn": 1}
    pri = {"content": 0, "structure": 1, "keywords": 2, "formatting": 3}
    todo = sorted([c for c in checks if c["status"] != "pass" and c["fix"]], key=lambda c: (order[c["status"]], pri[c["area"]]))
    return {"overall": overall, "scores": scores, "checks": checks, "top_fixes": [{"id": c["id"], "status": c["status"], "fix": c["fix"]} for c in todo[:3]],
            "counts": {"pass": sum(c["status"] == "pass" for c in checks), "warn": sum(c["status"] == "warn" for c in checks),
                       "fail": sum(c["status"] == "fail" for c in checks)},
            "basis": "Measured from the text extracted from your resume. Real applicant-tracking systems differ, so this is a check for common parsing "
                     "and clarity problems, not a score any employer uses."}


# =====================================================================================================
# 2. Coach: per-bullet diagnosis, honest prompts, and verification of any model-written rewrite
# =====================================================================================================
_VERB_OPTIONS = {"responsible for": ["Led", "Managed", "Owned", "Ran"], "worked on": ["Built", "Developed", "Implemented", "Improved"],
                 "helped": ["Contributed", "Supported", "Built", "Improved"], "involved in": ["Built", "Organised", "Contributed to"],
                 "assisted": ["Supported", "Built", "Tested"], "duties included": ["Built", "Managed", "Maintained"]}
SKELETON = "[Action verb] + [what you built or changed] + [tools you actually used] + [result: a number or a concrete outcome]"
_PLACEHOLDER = re.compile(r"\[[^\]\n]{1,60}\]")


def bullet_issues(text: str) -> List[Dict[str, Any]]:
    low = text.lower()
    wc = len(text.split())
    out: List[Dict[str, Any]] = []
    weak = [w for w in _WEAK if w in low]
    for w in weak:
        out.append({"id": "weak_phrase", "weight": 3, "message": f"'{w}' says you were present, not what you did.",
                    "question": "What exactly did you do? Name the action.", "verb_options": _VERB_OPTIONS.get(w, ["Built", "Improved"])})
    if _first_word(text) not in _VERBS and not weak:
        out.append({"id": "no_action_verb", "weight": 2, "message": "It does not start with an action verb.",
                    "question": "Start with the verb that is true: built, designed, tested, reduced, automated...", "verb_options": []})
    if not _METRIC.search(text):
        out.append({"id": "no_outcome", "weight": 3, "message": "No number or measurable outcome.",
                    "question": "How big was it or what changed (users, records, % faster, hours saved)? Use only a figure you can explain in an interview.",
                    "verb_options": []})
    if wc > 30:
        out.append({"id": "too_long", "weight": 2, "message": f"{wc} words is hard to scan.", "question": "Which one result matters most? Keep that and split or cut the rest.", "verb_options": []})
    if wc < 7:
        out.append({"id": "too_short", "weight": 2, "message": f"Only {wc} words.", "question": "What tool did you use and what was the result?", "verb_options": []})
    if _FIRST_PERSON.search(text):
        out.append({"id": "first_person", "weight": 1, "message": "First-person wording ('I', 'my').", "question": "Drop the pronoun and lead with the verb.", "verb_options": []})
    vague = [w for w in _VAGUE if re.search(r"\b" + re.escape(w) + r"\b", low)]
    if vague:
        out.append({"id": "vague_words", "weight": 1, "message": "Vague wording: " + ", ".join(vague) + ".", "question": "Replace it with the actual number or name.", "verb_options": []})
    buzz = [w for w in _BUZZ if w in low]
    if buzz:
        out.append({"id": "buzzwords", "weight": 1, "message": "Buzzword(s): " + ", ".join(buzz) + ".", "question": "Show it with evidence instead.", "verb_options": []})
    return out


def coach_bullets(text: str, limit: int = 5, only: Optional[str] = None) -> List[Dict[str, Any]]:
    """Weakest bullets first (or a single bullet the user pastes). Diagnosis only: no rewrite is made here."""
    if only is not None:
        items = [{"text": only.strip(), "section": ""}]
    else:
        items = parse_resume(text)["bullets"]
    ranked = []
    for i, b in enumerate(items):
        iss = bullet_issues(b["text"])
        ranked.append({"index": i, "section": b.get("section", ""), "bullet": b["text"], "issues": [{k: v for k, v in x.items() if k != "verb_options"} for x in iss],
                       "questions": list(dict.fromkeys(x["question"] for x in iss)),
                       "verb_options": list(dict.fromkeys(v for x in iss for v in x["verb_options"])),
                       "skeleton": SKELETON if iss else None, "severity": sum(x["weight"] for x in iss)})
    if only is None:
        ranked = [r for r in ranked if r["severity"] > 0]
        ranked.sort(key=lambda r: (-r["severity"], r["index"]))
    return ranked[:max(1, limit)]


_TOKEN_NEEDS_SOURCE = re.compile(r"[A-Za-z]*[A-Z][a-z]*[A-Z][A-Za-z]*|[A-Za-z]+[0-9][A-Za-z0-9]*|[A-Za-z]+[+#]+|[A-Za-z]+\.[A-Za-z]+")


def verify_rewrite(original: str, rewrite: str) -> Tuple[bool, List[str]]:
    """A rewrite may only rephrase. It must not add numbers, tools or names that the original bullet does not contain.
    Bracketed placeholders such as [add number] are allowed because they are visibly empty."""
    reasons: List[str] = []
    rw = _PLACEHOLDER.sub(" ", rewrite or "").strip()
    if not rw or len(rw.split()) < 4:
        return False, ["empty or too short"]
    if len(rw.split()) > 35:
        reasons.append("longer than 35 words")
    orig_low = (original or "").lower()
    orig_nums = {re.sub(r"[,\s]", "", n) for n in re.findall(r"\d[\d,]*(?:\.\d+)?", original or "")}
    for n in re.findall(r"\d[\d,]*(?:\.\d+)?", rw):
        if re.sub(r"[,\s]", "", n) not in orig_nums:
            reasons.append(f"new number '{n}' not in the original bullet")
    for tok in _TOKEN_NEEDS_SOURCE.findall(rw):
        if tok.lower() not in orig_low:
            reasons.append(f"new tool or name '{tok}' not in the original bullet")
    if _FIRST_PERSON.search(rw):
        reasons.append("uses first person")
    return (not reasons), reasons


# =====================================================================================================
# 3. Compare two resume versions
# =====================================================================================================
def compare_texts(a_text: str, b_text: str, a_label: str = "A", b_label: str = "B", skills_a: Optional[List[str]] = None,
                  skills_b: Optional[List[str]] = None, terms: Optional[List[str]] = None) -> Dict[str, Any]:
    pa, pb = parse_resume(a_text), parse_resume(b_text)
    ba, bb = [x["text"] for x in pa["bullets"]], [x["text"] for x in pb["bullets"]]
    used_b: set = set()
    identical: List[str] = []
    reworded: List[Dict[str, Any]] = []
    only_a: List[str] = []
    for ta in ba:
        best, best_s = None, 0.0
        for j, tb in enumerate(bb):
            if j in used_b:
                continue
            s = similarity(ta, tb)
            if s > best_s:
                best, best_s = j, s
        if best is not None and best_s >= 0.6:
            used_b.add(best)
            if _norm(ta) == _norm(bb[best]):
                identical.append(ta)
            else:
                reworded.append({"a": ta, "b": bb[best], "similarity": round(best_s, 2)})
        else:
            only_a.append(ta)
    only_b = [tb for j, tb in enumerate(bb) if j not in used_b]
    sa, sb = bullet_stats(pa["bullets"]), bullet_stats(pb["bullets"])
    ska = {s.lower(): s for s in (skills_a or [])}
    skb = {s.lower(): s for s in (skills_b or [])}
    secs_a, secs_b = {k for k in pa["sections"] if k != "header"}, {k for k in pb["sections"] if k != "header"}
    out: Dict[str, Any] = {
        "a": {"label": a_label, "words": len(a_text.split()), **sa}, "b": {"label": b_label, "words": len(b_text.split()), **sb},
        "bullets": {"identical": identical, "reworded": reworded, "only_a": only_a, "only_b": only_b},
        "skills": {"shared": sorted(ska.keys() & skb.keys()), "only_a": [ska[k] for k in sorted(ska.keys() - skb.keys())],
                   "only_b": [skb[k] for k in sorted(skb.keys() - ska.keys())]},
        "sections": {"shared": sorted(secs_a & secs_b), "only_a": sorted(secs_a - secs_b), "only_b": sorted(secs_b - secs_a)},
    }
    notes: List[str] = []
    for key, label in (("metric_ratio", "bullets with numbers"), ("verb_ratio", "bullets starting with an action verb")):
        d = round(100 * (sa[key] - sb[key]))
        if abs(d) >= 10:
            notes.append(f"{a_label if d > 0 else b_label} has more {label} ({abs(d)} points).")
    if abs(len(a_text.split()) - len(b_text.split())) >= 120:
        notes.append(f"{a_label if len(a_text.split()) > len(b_text.split()) else b_label} is noticeably longer.")
    if terms:
        la, lb = a_text.lower(), b_text.lower()
        ha = [t for t in terms if contains_term(la, t)]
        hb = [t for t in terms if contains_term(lb, t)]
        out["keywords"] = {"terms": terms, f"coverage_a": round(len(ha) / len(terms), 2), "coverage_b": round(len(hb) / len(terms), 2),
                           "only_a": [t for t in ha if t not in hb], "only_b": [t for t in hb if t not in ha],
                           "missing_in_both": [t for t in terms if t not in ha and t not in hb]}
        if len(ha) != len(hb):
            notes.append(f"{a_label if len(ha) > len(hb) else b_label} covers more of this posting's terms ({max(len(ha), len(hb))} vs {min(len(ha), len(hb))}).")
    out["notes"] = notes
    return out


# =====================================================================================================
# 4. Builder: validate a selection against the user's own text, suggest a draft, render exports
# =====================================================================================================
MAX_BUILD_BULLETS = 40
MAX_ITEM_CHARS = 400


def _own_index(own: Dict[str, Tuple[str, str]]) -> Tuple[List[Tuple[str, str, str]], List[Tuple[str, str, str]]]:
    """own = {resume_id: (label, raw_text)} -> (bullets, lines) as (text, resume_id, label)."""
    bullets, lines = [], []
    for rid, (label, text) in own.items():
        st = to_structure(text)
        for s in st["sections"]:
            for it in s["items"]:
                (bullets if it["type"] == "bullet" else lines).append((it["text"], rid, label))
    return bullets, lines


def validate_build(build: Dict[str, Any], own: Dict[str, Tuple[str, str]]) -> Tuple[Optional[Dict[str, Any]], List[str]]:
    """Returns (clean_build, errors). Every bullet, line and skill must come from one of the user's own resumes,
    or be explicitly marked by the user as written by them. Provenance is set here, never taken from the client."""
    errors: List[str] = []
    bullets, lines = _own_index(own)
    bnorm = {_norm(t): (rid, lab) for t, rid, lab in bullets}
    lnorm = {_norm(t): (rid, lab) for t, rid, lab in lines}
    all_low = " ".join(t.lower() for _, t in own.values())
    clean_secs: List[Dict[str, Any]] = []
    total = 0
    for si, s in enumerate(build.get("sections") or []):
        name = str(s.get("name") or "").strip()[:60]
        if not name:
            errors.append(f"Section {si + 1} has no name.")
            continue
        items = []
        for ii, it in enumerate(s.get("items") or []):
            text = re.sub(r"\s+", " ", str(it.get("text") or "")).strip()
            typ = "bullet" if it.get("type") == "bullet" else "line"
            where = f"{name} #{ii + 1}"
            if not text:
                continue
            if len(text) > MAX_ITEM_CHARS:
                errors.append(f"{where}: longer than {MAX_ITEM_CHARS} characters.")
                continue
            idx = bnorm if typ == "bullet" else lnorm
            hit = idx.get(_norm(text)) or (bnorm.get(_norm(text)) if typ == "line" else None)
            if hit:
                items.append({"type": typ, "text": text, "origin": "own", "source_resume_id": hit[0], "source_label": hit[1]})
            elif it.get("user_added") is True:
                items.append({"type": typ, "text": text, "origin": "user_added", "source_resume_id": None, "source_label": None})
            elif typ == "bullet" and it.get("edited") is True:
                near = max(((similarity(text, t), rid, lab) for t, rid, lab in bullets), default=(0.0, None, None))
                if near[0] >= 0.5:
                    items.append({"type": typ, "text": text, "origin": "edited", "source_resume_id": near[1], "source_label": near[2]})
                else:
                    errors.append(f"{where}: marked as edited but it does not resemble any bullet in your resumes. Mark it as written by you instead.")
            else:
                errors.append(f"{where}: not found in any of your resumes. If you wrote it yourself, mark it as added by you.")
            if typ == "bullet":
                total += 1
        clean_secs.append({"name": name, "key": name.lower(), "items": items})
    if total > MAX_BUILD_BULLETS:
        errors.append(f"Too many bullets ({total}). Keep it to {MAX_BUILD_BULLETS} or fewer.")
    skills_in = [re.sub(r"\s+", " ", str(x)).strip() for x in (build.get("skills") or []) if str(x).strip()]
    added = {re.sub(r"\s+", " ", str(x)).strip().lower() for x in (build.get("user_added_skills") or [])}
    skills: List[Dict[str, Any]] = []
    for sk in skills_in[:60]:
        if contains_term(all_low, sk.lower()):
            skills.append({"text": sk, "origin": "own"})
        elif sk.lower() in added:
            skills.append({"text": sk, "origin": "user_added"})
        else:
            errors.append(f"Skill '{sk}' does not appear in any of your resumes. If it is true for you, confirm it as added by you.")
    header = [re.sub(r"\s+", " ", str(h)).strip()[:120] for h in (build.get("header") or []) if str(h).strip()][:6]
    summary = re.sub(r"\s+", " ", str(build.get("summary") or "")).strip()[:400] or None
    if errors:
        return None, errors
    prov = Counter(it["origin"] for s in clean_secs for it in s["items"] if it["type"] == "bullet")
    prov.update(f"skill_{sk['origin']}" for sk in skills)
    return {"header": header, "summary": summary, "sections": clean_secs, "skills": skills, "provenance": dict(prov)}, []


def build_to_struct(build: Dict[str, Any]) -> Dict[str, Any]:
    """Renderer input: a build with provenance stripped. Skills become a final 'Skills' line section."""
    secs = [{"key": s["key"], "name": s["name"], "items": [{"type": it["type"], "text": it["text"]} for it in s["items"]]} for s in build["sections"]]
    if build.get("skills"):
        secs.append({"key": "skills", "name": "Skills", "items": [{"type": "line", "text": ", ".join(s["text"] for s in build["skills"])}]})
    return {"header": build.get("header") or [], "summary": build.get("summary"), "sections": secs}


def _bullet_score(text: str, terms: List[str]) -> float:
    low = text.lower()
    hits = sum(1 for t in terms if contains_term(low, t))
    return hits * 3 + (1.5 if _METRIC.search(text) else 0) + (1 if _first_word(text) in _VERBS else 0) - (1.5 if any(w in low for w in _WEAK) else 0)


def suggest_build(base_text: str, base_id: str, base_label: str, others: Dict[str, Tuple[str, str]], terms: List[str],
                  max_total: int = 14, per_block: int = 3) -> Dict[str, Any]:
    """A draft the user then edits. Orders and trims the BASE resume toward a posting. Adds nothing:
    bullets from other versions are only offered in could_add, never included."""
    st = to_structure(base_text)
    dropped: List[Dict[str, Any]] = []
    secs: List[Dict[str, Any]] = []
    skills_line: List[str] = []
    for s in st["sections"]:
        if s["key"] in _ALIASES["skills"]:
            for it in s["items"]:
                body = it["text"].split(":", 1)[1] if ":" in it["text"] else it["text"]
                skills_line += [x.strip() for x in re.split(r"[,;|/]", body) if 1 < len(x.strip()) <= 30]
            continue
        blocks: List[List[Dict[str, Any]]] = []
        for it in s["items"]:
            if it["type"] == "line" and s["key"] in _CONTENT:
                blocks.append([it])
            else:
                if not blocks:
                    blocks.append([])
                blocks[-1].append(it)
        kept_blocks = []
        for blk in blocks:
            bl = [(i, x) for i, x in enumerate(blk) if x["type"] == "bullet"]
            if len(bl) > per_block:
                ranked = sorted(bl, key=lambda p: (-_bullet_score(p[1]["text"], terms), p[0]))
                keep_idx = {p[0] for p in ranked[:per_block]}
                for i, x in bl:
                    if i not in keep_idx:
                        dropped.append({"bullet": x["text"], "section": s["name"], "reason": f"kept the {per_block} strongest bullets for this posting in this block"})
                blk = [x for i, x in enumerate(blk) if x["type"] != "bullet" or i in keep_idx]
            kept_blocks.append(blk)
        secs.append({"key": s["key"], "name": s["name"], "blocks": kept_blocks})

    def count(sx):
        return sum(1 for sec in sx for blk in sec["blocks"] for x in blk if x["type"] == "bullet")
    while count(secs) > max_total:   # drop the weakest remaining bullet overall
        cand = [(sec_i, blk_i, x_i, _bullet_score(x["text"], terms)) for sec_i, sec in enumerate(secs) if sec["key"] in _CONTENT
                for blk_i, blk in enumerate(sec["blocks"]) for x_i, x in enumerate(blk) if x["type"] == "bullet"]
        if not cand:
            break
        sec_i, blk_i, x_i, _ = min(cand, key=lambda c: (c[3], -c[0], -c[1], -c[2]))
        x = secs[sec_i]["blocks"][blk_i].pop(x_i)
        dropped.append({"bullet": x["text"], "section": secs[sec_i]["name"], "reason": "trimmed to fit one page"})
    out_secs = [{"key": sec["key"], "name": sec["name"], "items": [x for blk in sec["blocks"] for x in blk]} for sec in secs]
    skills_sorted = sorted(dict.fromkeys(skills_line), key=lambda sk: (0 if any(contains_term(sk.lower(), t) or contains_term(t, sk.lower()) for t in terms) else 1))
    chosen_low = " ".join(x["text"].lower() for sec in out_secs for x in sec["items"])
    could_add: List[Dict[str, Any]] = []
    obul, _ = _own_index({rid: v for rid, v in others.items() if rid != base_id})
    for t in terms:
        if contains_term(chosen_low, t) or contains_term(" ".join(skills_sorted).lower(), t):
            continue
        for text, rid, lab in obul:
            if contains_term(text.lower(), t):
                could_add.append({"term": t, "bullet": text, "source_resume_id": rid, "source_label": lab})
                break
        if len(could_add) >= 6:
            break
    present = [t for t in terms if contains_term(chosen_low + " " + " ".join(skills_sorted).lower(), t)]
    return {"base_resume_id": base_id, "base_label": base_label,
            "draft": {"header": st["header"][:6], "summary": None,
                      "sections": [{"name": s["name"], "items": [{"type": x["type"], "text": x["text"]} for x in s["items"]]} for s in out_secs],
                      "skills": skills_sorted[:30]},
            "dropped": dropped, "could_add": could_add, "terms_covered": present, "terms_missing": [t for t in terms if t not in present],
            "note": "This draft only reorders and trims the base resume. Nothing was added. Items under could_add come from your other versions and are not included until you add them."}


# ---- export
def safe_filename(label: str, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", (label or "resume").strip()).strip("._") or "resume"
    return f"{base[:60]}.{ext}"


_LATIN1_MAP = {"\u2022": "-", "\u25e6": "-", "\u25aa": "-", "\u25cf": "-", "\u25cb": "-", "\u2023": "-", "\u2013": "-", "\u2014": "-", "\u2018": "'",
               "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2026": "...", "\u2192": "->", "\u00a0": " "}


def _latin1(text: str, warnings: List[str]) -> str:
    out = []
    for ch in text:
        ch = _LATIN1_MAP.get(ch, ch)
        if ord(ch) > 255:
            warnings.append(f"'{ch}' cannot be shown in the PDF font")
            ch = "?"
        out.append(ch)
    return "".join(out)


def render_docx(struct: Dict[str, Any]) -> bytes:
    """Single-column, real-text Word file: standard headings, real bullet style, no tables, no images."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    for i, ln in enumerate(struct.get("header") or []):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(ln)
        r.bold = i == 0
        r.font.size = Pt(16 if i == 0 else 10)

    def heading(name: str):
        p = d.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(3)
        r = p.add_run(name.upper())
        r.bold, r.font.size = True, Pt(11)
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        bt = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "444444")):
            bt.set(qn(k), v)
        b.append(bt)
        pPr.append(b)
    if struct.get("summary"):
        heading("Summary")
        d.add_paragraph(struct["summary"])
    for s in struct.get("sections") or []:
        heading(s.get("name") or "")
        for it in s.get("items") or []:
            if it.get("type") == "bullet":
                d.add_paragraph(it["text"], style="List Bullet")
            else:
                p = d.add_paragraph()
                p.add_run(it["text"]).bold = (s.get("key") in _CONTENT)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_pdf(struct: Dict[str, Any]) -> Tuple[bytes, List[str]]:
    """Single-column PDF with real text (copy-pasteable and parseable). Returns (bytes, warnings)."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as e:
        raise RuntimeError("PDF export needs the 'reportlab' package (pip install reportlab).") from e
    from xml.sax.saxutils import escape
    warns: List[str] = []

    def t(x: str) -> str:
        return escape(_latin1(x, warns))
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=9.5, leading=12.2)
    name_st = ParagraphStyle("n", parent=base, fontName="Helvetica-Bold", fontSize=16, leading=19, alignment=1)
    head_st = ParagraphStyle("h", parent=base, fontSize=9.5, leading=12, alignment=1)
    sec_st = ParagraphStyle("s", parent=base, fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=7, spaceAfter=1)
    blt_st = ParagraphStyle("bl", parent=base, leftIndent=11, bulletIndent=2, spaceAfter=1.2)
    bold_st = ParagraphStyle("bd", parent=base, fontName="Helvetica-Bold", spaceBefore=2)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title="Resume")
    flow: List[Any] = []
    for i, ln in enumerate(struct.get("header") or []):
        flow.append(Paragraph(t(ln), name_st if i == 0 else head_st))

    def section(name: str):
        flow.append(Paragraph(t(name.upper()), sec_st))
        flow.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#444444"), spaceAfter=3))
    if struct.get("summary"):
        section("Summary")
        flow.append(Paragraph(t(struct["summary"]), base))
    for s in struct.get("sections") or []:
        section(s.get("name") or "")
        for it in s.get("items") or []:
            if it.get("type") == "bullet":
                flow.append(Paragraph(t(it["text"]), blt_st, bulletText="-"))
            else:
                flow.append(Paragraph(t(it["text"]), bold_st if s.get("key") in _CONTENT else base))
    if not flow:
        flow.append(Spacer(1, 1))
    doc.build(flow)
    return buf.getvalue(), sorted(set(warns))
