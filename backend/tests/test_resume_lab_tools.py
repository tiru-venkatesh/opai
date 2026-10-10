"""Run: cd backend && python -m pytest tests/test_resume_lab_tools.py -q
Pure functions only: no server, no database, no network."""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

import resume_lab_tools as T  # noqa: E402

FULL = """Asha Rao
asha@example.com | +91 98765 43210 | github.com/asha | linkedin.com/in/asha
EDUCATION
B.Tech Computer Science, JNTU Narasaraopeta, 2024-2028
Projects
OPAI | Python, FastAPI, RAG
- Built a RAG assistant in Python and FastAPI that answers questions over 500 documents with 92% retrieval accuracy
- Fine-tuned a PyTorch classifier on 20k samples, improving F1 by 8%
- Responsible for various tasks in the team
CivicIQ | React, Node.js
- Designed a React dashboard with Node.js that reduced triage time by 40%
- Worked on the website frontend
Experience
- Developed REST APIs with Docker and deployed them on AWS serving 1,000+ requests per day
Skills
Languages: Python, JavaScript, SQL
Tools: Docker, PyTorch, React
"""
WEB = """Asha Rao
asha@example.com | +91 98765 43210
Projects
- Designed a React dashboard with TypeScript that reduced report time by 40%
- Built a Node API with 12 endpoints and automated tests
Skills
React, TypeScript, Node
"""


def checks(r):
    return {c["id"]: c for c in r["checks"]}


# ---------------------------------------------------------------- structure
def test_structure_keeps_order_and_types():
    st = T.to_structure(FULL)
    assert st["header"][0] == "Asha Rao" and len(st["header"]) == 2
    assert [s["key"] for s in st["sections"]] == ["education", "projects", "experience", "skills"]
    proj = [s for s in st["sections"] if s["key"] == "projects"][0]["items"]
    assert proj[0] == {"type": "line", "text": "OPAI | Python, FastAPI, RAG"}
    assert proj[1]["type"] == "bullet" and proj[1]["text"].startswith("Built a RAG assistant")
    assert [i["type"] for i in proj] == ["line", "bullet", "bullet", "bullet", "line", "bullet", "bullet"]


def test_structure_agrees_with_parse_resume_on_bullets():
    a = [b["text"] for b in T.parse_resume(FULL)["bullets"]]
    b = [i["text"] for s in T.to_structure(FULL)["sections"] for i in s["items"] if i["type"] == "bullet" and s["key"] in T._CONTENT]
    assert a == b


def test_render_text_round_trips_content():
    out = T.render_text(T.to_structure(FULL))
    assert "PROJECTS" in out and "- Built a RAG assistant" in out and "asha@example.com" in out


def test_wrapped_lines_merge_but_project_titles_do_not():
    txt = "Projects\n- Improved latency by 35% through caching and\nbatching of requests\nCivicIQ | React, Node.js\n- Built a triage dashboard used by 40 staff members daily\nIntern, Acme Corp  Jun 2025 - Aug 2025\n- Wrote tests that caught 12 regressions before release\n"
    b = [x["text"] for x in T.parse_resume(txt)["bullets"]]
    assert b[0] == "Improved latency by 35% through caching and batching of requests"
    assert len(b) == 3 and all("CivicIQ" not in x and "Acme" not in x for x in b)


# ---------------------------------------------------------------- review
def test_review_good_resume_passes_core_checks():
    r = T.review_resume(FULL)
    c = checks(r)
    assert c["core_sections"]["status"] == "pass" and c["contact"]["status"] == "pass" and c["profile_links"]["status"] == "pass"
    assert c["weak_phrases"]["status"] == "warn" and any("Responsible" in e for e in c["weak_phrases"]["evidence"])
    assert 0 <= r["overall"] <= 100 and set(r["scores"]) == {"structure", "formatting", "content"}
    assert "keyword_coverage" not in c        # no posting supplied


def test_review_flags_missing_contact_and_sections():
    r = T.review_resume("Asha\nPassionate about tech.\nHobbies\n- Chess\n")
    c = checks(r)
    assert c["contact"]["status"] == "fail" and c["core_sections"]["status"] == "fail"
    assert c["bullet_count"]["status"] == "fail" and r["top_fixes"]


def test_review_detects_spacing_odd_chars_dates_and_glyphs():
    txt = ("Asha Rao asha@x.com +91 98765 43210\nP R O J E C T S\nT e n s o r F l o w  P y T o r c h\n"
           "Projects\n• Built a thing with Python that served 100 users daily\n- Designed an API used by 20 students every week\n"
           "★ rated app\nJan 2024 - 03/2024\n")
    c = checks(T.review_resume(txt))
    assert c["spaced_text"]["status"] == "warn" and c["special_characters"]["status"] == "warn"
    assert c["date_format"]["status"] == "warn" and c["bullet_glyphs"]["status"] == "warn"


def test_review_keywords_and_stuffing():
    terms = ["python", "pytorch", "kubernetes", "react"]
    c = checks(T.review_resume(FULL, terms))
    assert c["keyword_coverage"]["status"] == "pass"       # python, pytorch, react present = 3 of 4
    assert "kubernetes" in c["keyword_coverage"]["evidence"] and "python" not in c["keyword_coverage"]["evidence"]
    stuffed = FULL + "\n" + "\n".join(["- Used Python daily for work"] * 6)
    assert checks(T.review_resume(stuffed, ["python"]))["keyword_stuffing"]["status"] == "warn"


def test_review_length_check():
    long_txt = FULL + "\n- " + " ".join(["word"] * 1000)
    assert checks(T.review_resume(long_txt))["length"]["status"] == "fail"
    assert checks(T.review_resume(FULL))["length"]["status"] == "pass"


# ---------------------------------------------------------------- coach
def test_coach_ranks_weakest_first_and_asks_questions_not_invents():
    out = T.coach_bullets(FULL, limit=3)
    assert out[0]["bullet"] == "Responsible for various tasks in the team"
    ids = {i["id"] for i in out[0]["issues"]}
    assert {"weak_phrase", "no_outcome", "vague_words"} <= ids
    assert "Led" in out[0]["verb_options"] and out[0]["questions"] and out[0]["skeleton"]
    assert all(r["severity"] > 0 for r in out)
    assert not any("rewrite" in r for r in out)


def test_coach_single_bullet_and_strong_bullet_has_no_issues():
    one = T.coach_bullets("", only="Built a RAG assistant in Python that answers questions over 500 documents with 92% retrieval accuracy")
    assert one[0]["issues"] == [] and one[0]["skeleton"] is None


def test_verify_rewrite_blocks_invented_facts():
    orig = "Built a RAG assistant in Python over 500 documents"
    ok, why = T.verify_rewrite(orig, "Built a Python RAG assistant that searches 500 documents")
    assert ok and not why
    ok, why = T.verify_rewrite(orig, "Built a RAG assistant in Python over 500 documents, cutting latency by 40%")
    assert not ok and any("40" in w for w in why)
    ok, why = T.verify_rewrite(orig, "Built a RAG assistant using LangChain over 500 documents")
    assert not ok and any("LangChain" in w for w in why)
    ok, why = T.verify_rewrite(orig, "Built a Python RAG assistant over 500 documents, improving accuracy by [add number]%")
    assert ok, why
    assert not T.verify_rewrite(orig, "I built a RAG assistant in Python for 500 documents")[0]
    assert not T.verify_rewrite(orig, "ok")[0]


# ---------------------------------------------------------------- compare
def test_compare_aligns_bullets_skills_and_terms():
    a = FULL
    b = FULL.replace("Fine-tuned a PyTorch classifier on 20k samples, improving F1 by 8%",
                     "Fine-tuned a PyTorch classifier on 20k samples and improved F1 by 8%").replace(
        "Skills\nLanguages", "- Automated deployment checks across 12 services using Docker\nSkills\nLanguages")
    r = T.compare_texts(a, b, "A", "B", ["Python", "Docker"], ["Python", "PyTorch"], terms=["python", "docker", "kubernetes"])
    assert len(r["bullets"]["reworded"]) == 1 and r["bullets"]["reworded"][0]["similarity"] >= 0.6
    assert any("Automated deployment" in x for x in r["bullets"]["only_b"]) and r["bullets"]["only_a"] == []
    assert r["skills"]["only_a"] == ["Docker"] and r["skills"]["only_b"] == ["PyTorch"] and r["skills"]["shared"] == ["python"]
    assert r["keywords"]["missing_in_both"] == ["kubernetes"]


def test_compare_notes_name_the_stronger_version():
    r = T.compare_texts(FULL, WEB.replace("reduced report time by 40%", "reduced report time"), "ML", "Web")
    assert r["a"]["label"] == "ML" and r["b"]["bullets"] == 2
    assert any("ML has more bullets with numbers" in n for n in r["notes"])


# ---------------------------------------------------------------- build
OWN = {"r1": ("ML", FULL), "r2": ("Web", WEB)}


def _build(items, **kw):
    return {"header": ["Asha Rao", "asha@example.com"], "sections": [{"name": "Projects", "items": items}], "skills": ["Python", "React"], **kw}


def test_build_accepts_own_bullets_and_sets_provenance_server_side():
    b = _build([{"type": "bullet", "text": "Built a Node API with 12 endpoints and automated tests", "source_resume_id": "FAKE"}])
    clean, err = T.validate_build(b, OWN)
    assert not err
    it = clean["sections"][0]["items"][0]
    assert it["origin"] == "own" and it["source_resume_id"] == "r2" and it["source_label"] == "Web"      # client value ignored
    assert clean["provenance"]["own"] == 1


def test_build_rejects_invented_bullets_and_skills():
    clean, err = T.validate_build(_build([{"type": "bullet", "text": "Led a team of 30 engineers at Google"}], skills=["Python", "Kubernetes"]), OWN)
    assert clean is None and any("not found in any of your resumes" in e for e in err) and any("Kubernetes" in e for e in err)


def test_build_allows_user_marked_additions_and_edits():
    items = [{"type": "bullet", "text": "Won the college hackathon with a team of four", "user_added": True},
             {"type": "bullet", "text": "Designed a React dashboard with TypeScript that cut report time by 40%", "edited": True},
             {"type": "bullet", "text": "Completely unrelated sentence about cooking dinner tonight", "edited": True}]
    clean, err = T.validate_build(_build(items), OWN)
    assert clean is None and len(err) == 1 and "does not resemble" in err[0]
    clean, err = T.validate_build(_build(items[:2], skills=["Python", "Rust"], user_added_skills=["rust"]), OWN)
    assert not err
    assert [i["origin"] for i in clean["sections"][0]["items"]] == ["user_added", "edited"]
    assert clean["provenance"]["skill_user_added"] == 1


def test_build_limits():
    many = [{"type": "bullet", "text": f"Built a Node API with 12 endpoints and automated tests", "user_added": True}] * 41
    assert T.validate_build(_build(many), OWN)[0] is None
    assert T.validate_build(_build([{"type": "bullet", "text": "x" * 401, "user_added": True}]), OWN)[0] is None


def test_suggest_build_trims_orders_and_adds_nothing():
    terms = ["react", "node.js", "kubernetes", "typescript"]
    r = T.suggest_build(FULL, "r1", "ML", OWN, terms, max_total=3, per_block=2)
    d = r["draft"]
    bullets = [i["text"] for s in d["sections"] for i in s["items"] if i["type"] == "bullet"]
    own_bullets = {b["text"] for b in T.parse_resume(FULL)["bullets"]}
    assert len(bullets) <= 3 and set(bullets) <= own_bullets                  # selection only
    assert r["dropped"] and all("reason" in x for x in r["dropped"])
    assert "Python" in d["skills"] and d["skills"].index("React") < d["skills"].index("SQL")   # job-relevant skill first
    assert any(x["term"] == "typescript" and "Web" == x["source_label"] for x in r["could_add"])
    assert "typescript" in r["terms_missing"] and "kubernetes" in r["terms_missing"]
    clean, err = T.validate_build(d, OWN)                                      # the draft always validates
    assert not err, err


# ---------------------------------------------------------------- export
def _built():
    clean, err = T.validate_build(_build([{"type": "line", "text": "OPAI | Python, FastAPI, RAG"},
                                          {"type": "bullet", "text": "Fine-tuned a PyTorch classifier on 20k samples, improving F1 by 8%"}]), OWN)
    assert not err, err
    return T.build_to_struct(clean)


def test_render_docx_is_single_column_text():
    from docx import Document
    data = T.render_docx(_built())
    d = Document(io.BytesIO(data))
    assert len(d.tables) == 0 and not d.inline_shapes
    texts = [p.text for p in d.paragraphs]
    assert texts[0] == "Asha Rao" and "PROJECTS" in texts and "SKILLS" in texts
    assert any(p.style.name == "List Bullet" and p.text.startswith("Fine-tuned") for p in d.paragraphs)


def test_render_pdf_has_real_extractable_text():
    data, warns = T.render_pdf(_built())
    assert data[:5] == b"%PDF-" and warns == []
    from pypdf import PdfReader
    txt = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
    assert "Asha Rao" in txt and "PROJECTS" in txt and "Fine-tuned a PyTorch classifier" in txt
    back = T.parse_resume(txt)                       # a parser can read our own export
    assert any("Fine-tuned" in b["text"] for b in back["bullets"])


def test_render_pdf_warns_on_unrenderable_characters():
    st = {"header": ["Veera \u0c35\u0c40\u0c30"], "sections": []}
    data, warns = T.render_pdf(st)
    assert data[:5] == b"%PDF-" and warns


def test_safe_filename():
    assert T.safe_filename("ML Intern / v2", "pdf") == "ML_Intern_v2.pdf"
    assert "/" not in T.safe_filename("../../etc/passwd", "txt") and T.safe_filename("", "txt") == "resume.txt"


if __name__ == "__main__":
    fails = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print("PASS", name)
            except Exception as e:  # noqa: BLE001
                fails += 1
                import traceback
                print("FAIL", name, "->", repr(e))
                traceback.print_exc(limit=3)
    print("failures:", fails)
    sys.exit(1 if fails else 0)
